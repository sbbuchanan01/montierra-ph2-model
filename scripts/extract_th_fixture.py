import openpyxl, json, datetime
wb = openpyxl.load_workbook('th.xlsx', data_only=True)
def v(ws, addr):
    x = ws[addr].value
    if isinstance(x, datetime.datetime): return x.strftime('%Y-%m-%d')
    return x
def col(ws, c, r0, r1):
    return [v(ws, f'{c}{r}') for r in range(r0, r1+1)]
def rowvals(ws, r, c0, c1):
    from openpyxl.utils import get_column_letter as L, column_index_from_string as I
    return [v(ws, f'{L(c)}{r}') for c in range(I(c0), I(c1)+1)]
inp = wb['Inputs']; um = wb['Unit Mix']; us = wb['Unit Schedule']; db = wb['Development Budget']
mc = wb['Margin on Cost']; dr = wb['Draw Schedule']; ds = wb['Debt Schedule']; cf = wb['Sales CF']
an = wb['Annual Summary']; wf = wb['Waterfall']; rt = wb['Returns']; tx = wb['Taxes']
out = {}
out['inputs'] = {a: v(inp, 'F'+str(r)) for a, r in {
 'units':10,'nsf':12,'avgSf':13,'gsf':15,'lastDelivery':30,'sellout':34,'selloutDate':35,'selloutDur':36,
 'waPrice':40,'waPsf':41,'grossToday':42,'totalDeductions':48,'sizingRate':96,'totalCost':103,'costExFee':104,
 'lhs':105,'maxLtc':106,'maxLtgs':107,'commitment':108,'binding':109,'commitPctCost':110,'commitPctSellout':111,
 'loanPerUnit':112,'equity':113,'gpEquity':126,'lpEquity':127}.items()}
out['carry'] = {r: [v(inp,f'E{r}'), v(inp,f'F{r}'), v(inp,f'G{r}')] for r in range(64,72)}
out['indices'] = {'price': rowvals(inp,58,'F','O'), 'carry': rowvals(inp,59,'F','O'), 'tax': rowvals(inp,60,'F','O')}
out['unitMixTotals'] = {c: v(um, f'{c}18') for c in 'FGHIJKLM'}
out['unitSchedule'] = [{c: v(us, f'{c}{r}') for c in 'BCDEFGHIJKLMN'} for r in range(5, 15)]
out['unitScheduleTotals'] = {'D':v(us,'D505'),'E':v(us,'E505'),'F':v(us,'F505'),'M':v(us,'M505'),'N':v(us,'N505'),'H507':v(us,'H507'),'K507':v(us,'K507')}
out['budget'] = {str(r): {'E':v(db,f'E{r}'),'F':v(db,f'F{r}'),'G':v(db,f'G{r}'),'H':v(db,f'H{r}')} for r in list(range(8,14))+list(range(16,26))+list(range(28,39))+list(range(41,47))+list(range(49,53))+[54]}
out['budgetSources'] = {str(r): v(db,f'E{r}') for r in range(59,65)}
out['budgetMetrics'] = {str(r): v(db,f'E{r}') for r in range(67,84)}
out['budgetChecks'] = {str(r): v(db,f'E{r}') for r in [86,87,88,89]}
out['margin'] = {str(r): {'D':v(mc,f'D{r}'),'F':v(mc,f'F{r}')} for r in list(range(4,5))+list(range(7,14))+list(range(15,19))+list(range(20,24))+list(range(25,31))+[32]}
out['taxes'] = {str(r): rowvals(tx, r, 'C','M') for r in [11,12,13,14,15,16,17,20,21,22,23,24]}
out['draw'] = {c: col(dr, c, 11, 131) for c in 'CEFGHIJKLMNOPQRST'}
out['debtHead'] = {'F4':v(ds,'F4'),'F5':v(ds,'F5'),'F6':v(ds,'F6'),'F7':v(ds,'F7'),'F8':v(ds,'F8'),'G8':v(ds,'G8')}
out['debt'] = {c: col(ds, c, 11, 131) for c in 'FGHIJKLMNO'}
cfrows = [6,7,8,9,10,11,12,13,14,15,16]+list(range(19,39))+list(range(41,50))+[51,53]+list(range(56,63))+[64,70,71,72,73,74,75,76,78,82,85]
out['salesCf'] = {str(r): rowvals(cf, r, 'D','DT') for r in cfrows}
out['salesCfTotals'] = {str(r): v(cf,f'C{r}') for r in cfrows+[65,66,67,79,80,81]}
out['annual'] = {str(r): rowvals(an, r, 'C','N') for r in [7,8,9,10,11,12,13,14,15,16,17,18,19,22,23,24,25,26,29,30,31,32,33,35,38,39,40,41,42,43,45]}
out['waterfall'] = {str(r): rowvals(wf, r, 'C','DS') for r in list(range(9,11))+list(range(12,14))+list(range(15,33))+list(range(34,40))+list(range(41,47))}
out['waterfallSummary'] = {str(r): [v(wf,f'C{r}'), v(wf,f'D{r}')] for r in range(52,56)}
out['returns'] = {}
for r in range(6,24):
    for c in 'DJ':
        x = v(rt, f'{c}{r}')
        if x is not None: out['returns'][f'{c}{r}'] = x
for r in [31,32,33,44,45,46,47,48,49,63,64,65,67,68,69]:
    out['returns'][f'row{r}'] = rowvals(rt, r, 'C','M')
for r in [35,36,37,38,51,52,53,54,55,56,57,72,73,74,75,76,77,78,79,80,81,82,83]:
    out['returns'][f'F{r}'] = v(rt, f'F{r}')
out['sensProfit'] = [rowvals(rt, r, 'D','L') for r in range(90,99)]
out['sensMargin'] = [rowvals(rt, r, 'D','L') for r in range(103,112)]
out['dashboard'] = {a: v(wb['Dashboard'], a) for a in ['K9','K10','K11','K12','K13','K14','K17','K19','K21','K22','K23','K24','K25','K29','D21','D22','D23','D24','D26']}
json.dump(out, open('th-fixture.json','w'), indent=0, default=str)
print(len(json.dumps(out, default=str)))
print(out['inputs'])
print(out['debtHead'])
